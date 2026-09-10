"""
mix_log.py
----------
Keeps two separate Excel logs of the work done in the mix design tool.

02 MixLog.xlsx      sheet "MixDesign"
    One row per 1 m3 mix design.
    Columns : Timestamp | <material> = the proportion you entered | Volume (m3) = 1
              | Total mass (kg/m3) | Cost ($/m3) | Energy (MJ/m3) | CO2 (kg/m3)

03 ActualBatch.xlsx sheet "ActualBatch"
    One row per batch actually cast.
    Columns : Timestamp | <material> = mass in kg | Volume (m3)
              | Total mass (kg) | Cost ($) | Energy (MJ) | CO2 (kg)

A material used for the first time in a later run becomes a new column; earlier
rows are filled with 0. Summary columns always stay at the right-hand end.

Quick start:
    from mix_log import save_mix_log, save_batch_log

    save_mix_log(binder_ratios, fiber_ratios_percent, mix_result)
    save_batch_log(actual_result, 0.05)
"""

from __future__ import annotations

import os
import re
import shutil
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Any, Mapping, Optional, Sequence

import pandas as pd

# ------------------------------------------------------------- settings ----

MIX_LOG_FILE = "02 MixLog.xlsx"        # one row per 1 m3 design
MIX_SHEET_NAME = "MixDesign"

BATCH_LOG_FILE = "03 ActualBatch.xlsx"  # one row per cast batch
BATCH_SHEET_NAME = "ActualBatch"

TIME_COLUMN = "Timestamp"
TIME_FORMAT = "%Y-%m-%d %H:%M:%S"

# Suffix appended to fiber names in the mix design log, where fibers are given
# as percent of volume while binders are mass ratios. Set to "" for plain names.
FIBER_COLUMN_SUFFIX = " (%vol)"

# Summary columns, always kept at the right-hand end of each sheet
MIX_TRAILING = (
    "Volume (m3)",
    "Total mass (kg/m3)",
    "Cost ($/m3)",
    "Energy (MJ/m3)",
    "CO2 (kg/m3)",
)
BATCH_TRAILING = (
    "Volume (m3)",
    "Total mass (kg)",
    "Cost ($)",
    "Energy (MJ)",
    "CO2 (kg)",
)

# Materials unused in a given run -> write 0 (True) or leave blank (False)
FILL_MISSING_WITH_ZERO = True

# Backwards-compatible alias, some older code imports this name
LOG_EXCEL_FILE = MIX_LOG_FILE


# -------------------------------------------------------------- helpers ----

def _clean_name(value: Any) -> str:
    """
    Normalizes a material name into a single-line column header.

    Names in the Excel database contain line breaks, e.g.
    "OPC\\n(Ordinary Portland Cement)". Collapsing them keeps the log headers
    on one line and identical to the ones used by the standalone HTML version.
    """
    text = str(value).replace("\xa0", " ")
    return re.sub(r"\s+", " ", text).strip()


def _order_columns(old_columns, new_row: Mapping[str, Any],
                   trailing: Sequence[str]) -> list[str]:
    """Keeps existing column order, appends new materials, pushes summaries last."""
    ordered = [TIME_COLUMN]
    ordered += [c for c in old_columns if c not in ordered and c not in trailing]
    ordered += [c for c in new_row if c not in ordered and c not in trailing]
    ordered += [c for c in trailing if c in old_columns or c in new_row]
    return ordered


def _format_sheet(path: Path, sheet_name: str) -> None:
    """Styles the sheet: bold header, frozen header row, sensible column widths."""
    try:
        from openpyxl import load_workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        return

    wb = load_workbook(path)
    if sheet_name not in wb.sheetnames:
        return
    ws = wb[sheet_name]

    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.freeze_panes = "B2"

    for col_idx, column_cells in enumerate(ws.columns, start=1):
        letter = column_cells[0].column_letter
        max_len = max((len(str(c.value)) for c in column_cells if c.value is not None),
                      default=8)
        ws.column_dimensions[letter].width = min(max(10, max_len + 2), 24)
        if col_idx > 1:
            for cell in column_cells[1:]:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "0.###"
                    cell.alignment = Alignment(horizontal="center")

    wb.save(path)


def _read_log(path: str | Path, sheet_name: str) -> pd.DataFrame:
    """Reads a log workbook. Returns an empty DataFrame if it does not exist yet."""
    path = Path(path)
    if not path.exists():
        return pd.DataFrame(columns=[TIME_COLUMN])
    try:
        return pd.read_excel(path, sheet_name=sheet_name)
    except (ValueError, KeyError):
        return pd.DataFrame(columns=[TIME_COLUMN])


def _append_row(row: Mapping[str, Any], path: str | Path, sheet_name: str,
                trailing: Sequence[str]) -> Path:
    """Appends one row to a log workbook, widening the table if needed."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    old_df = _read_log(path, sheet_name)
    new_df = pd.concat([old_df, pd.DataFrame([row])], ignore_index=True)
    new_df = new_df.reindex(columns=_order_columns(list(old_df.columns), row, trailing))

    if FILL_MISSING_WITH_ZERO:
        for col in new_df.columns:
            if col == TIME_COLUMN:
                continue
            new_df[col] = pd.to_numeric(new_df[col], errors="coerce").fillna(0.0)

    # Write to a temporary file first, then replace -> the existing log is never
    # corrupted if something fails halfway through.
    tmp_fd, tmp_name = tempfile.mkstemp(suffix=".xlsx", dir=str(path.parent))
    os.close(tmp_fd)
    tmp_path = Path(tmp_name)
    try:
        with pd.ExcelWriter(tmp_path, engine="openpyxl") as writer:
            new_df.to_excel(writer, sheet_name=sheet_name, index=False)
        _format_sheet(tmp_path, sheet_name)
        shutil.move(str(tmp_path), str(path))
    except PermissionError as exc:
        tmp_path.unlink(missing_ok=True)
        raise PermissionError(
            f"Cannot write to '{path}'. Please close the file in Excel and try again."
        ) from exc
    except Exception:
        tmp_path.unlink(missing_ok=True)
        raise

    return path


def _stamp(timestamp: Optional[datetime]) -> str:
    return (timestamp or datetime.now()).strftime(TIME_FORMAT)


# --------------------------------------------------------- main functions ----

def read_mix_log(path: str | Path = MIX_LOG_FILE) -> pd.DataFrame:
    """Reads the whole mix design history."""
    return _read_log(path, MIX_SHEET_NAME)


def read_batch_log(path: str | Path = BATCH_LOG_FILE) -> pd.DataFrame:
    """Reads the whole actual-batch history."""
    return _read_log(path, BATCH_SHEET_NAME)


def save_mix_log(
    binder_ratios: Optional[Mapping[str, float]] = None,
    fiber_ratios_percent: Optional[Mapping[str, float]] = None,
    mix_result: Optional[dict] = None,
    path: str | Path = MIX_LOG_FILE,
    timestamp: Optional[datetime] = None,
) -> Path:
    """
    Appends one row to 02 MixLog.xlsx describing a 1 m3 mix design.

    binder_ratios        : dict {material name: ratio} - exactly the calculation input
    fiber_ratios_percent : dict {fiber name: volume percent}
    mix_result           : output of calculate_mix_design(), for the summary columns
    """
    if not binder_ratios and not fiber_ratios_percent:
        raise ValueError("Nothing to log: at least one material is required.")

    row: dict[str, Any] = {TIME_COLUMN: _stamp(timestamp)}

    for name, ratio in (binder_ratios or {}).items():
        row[_clean_name(name)] = float(ratio)
    for name, percent in (fiber_ratios_percent or {}).items():
        row[f"{_clean_name(name)}{FIBER_COLUMN_SUFFIX}"] = float(percent)

    row["Volume (m3)"] = 1.0
    totals = (mix_result or {}).get("totals_per_m3", {})
    row["Total mass (kg/m3)"] = round(float(totals.get("total_material_mass_kg_m3", 0.0)), 2)
    row["Cost ($/m3)"] = round(float(totals.get("cost_per_m3", 0.0)), 3)
    row["Energy (MJ/m3)"] = round(float(totals.get("energy_mj_per_m3", 0.0)), 2)
    row["CO2 (kg/m3)"] = round(float(totals.get("co2_kg_per_m3", 0.0)), 3)

    return _append_row(row, path, MIX_SHEET_NAME, MIX_TRAILING)


def save_batch_log(
    actual_result: dict,
    actual_volume_m3: Optional[float] = None,
    path: str | Path = BATCH_LOG_FILE,
    timestamp: Optional[datetime] = None,
) -> Path:
    """
    Appends one row to 03 ActualBatch.xlsx: the mass of each material, in kilograms,
    for the volume actually cast.

    actual_result    : output of calculate_actual_volume()
    actual_volume_m3 : batch volume; taken from actual_result when omitted
    """
    if not actual_result:
        raise ValueError("Nothing to log: no batch result was given.")

    volume = actual_volume_m3
    if volume is None:
        volume = actual_result.get("actual_volume_m3", 0.0)

    row: dict[str, Any] = {TIME_COLUMN: _stamp(timestamp)}

    for group in ("binder", "fiber"):
        for item in actual_result.get(group, {}).values():
            row[_clean_name(item["name"])] = round(float(item["mass_kg"]), 3)

    totals = actual_result.get("totals", {})
    row["Volume (m3)"] = float(volume)
    row["Total mass (kg)"] = round(float(totals.get("total_material_mass_kg", 0.0)), 3)
    row["Cost ($)"] = round(float(totals.get("cost", 0.0)), 3)
    row["Energy (MJ)"] = round(float(totals.get("energy_mj", 0.0)), 3)
    row["CO2 (kg)"] = round(float(totals.get("co2_kg", 0.0)), 3)

    return _append_row(row, path, BATCH_SHEET_NAME, BATCH_TRAILING)
